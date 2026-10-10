"""Bounded local import/export; never fetches remote resources."""
import io
import re
import zipfile
from pathlib import PurePosixPath
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils.html import escape
from .models import ManualNode
from .policy import require, node_allowed, require_node
from .services import render_markdown, save_revision, audit

MAX_UPLOAD = 10 * 1024 * 1024
MAX_EXPANDED = 20 * 1024 * 1024
MAX_ENTRIES = 100
SUPPORTED = {'.md', '.html', '.htm', '.docx'}

def validate_archive(data):
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        entries = archive.infolist()
        if len(entries) > MAX_ENTRIES:
            raise ValidationError('ZIP 항목은 최대 100개입니다.')
        total = 0
        for info in entries:
            path = PurePosixPath(info.filename)
            if path.is_absolute() or '..' in path.parts or '\\' in info.filename or ':' in info.filename:
                raise ValidationError('안전하지 않은 ZIP 경로입니다.')
            if (info.external_attr >> 16) & 0o170000 == 0o120000 or info.flag_bits & 1:
                raise ValidationError('링크 또는 암호화된 ZIP은 허용하지 않습니다.')
            total += info.file_size
            if total > MAX_EXPANDED or info.file_size > MAX_UPLOAD:
                raise ValidationError('ZIP 압축 해제 크기 제한을 초과했습니다.')
            if info.file_size > max(info.compress_size, 1) * 100:
                raise ValidationError('압축률 제한을 초과했습니다.')
        return archive
    except (zipfile.BadZipFile, ValueError) as exc:
        raise ValidationError('올바른 ZIP 파일이 아닙니다.') from exc

def decode_document(name, data):
    suffix = PurePosixPath(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValidationError('Markdown, HTML, DOCX 파일만 가져올 수 있습니다.')
    if len(data) > MAX_UPLOAD:
        raise ValidationError('파일 크기는 10MB 이하입니다.')
    if suffix == '.docx':
        from docx import Document
        archive = validate_archive(data)
        try:
            if 'word/document.xml' not in archive.namelist():
                raise ValidationError('올바른 DOCX 문서가 아닙니다.')
        finally:
            archive.close()
        try:
            doc = Document(io.BytesIO(data))
            blocks = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.replace('<', '&lt;').replace('>', '&gt;')
                if not text.strip():
                    continue
                style = paragraph.style.name if paragraph.style else ''
                heading = re.search(r'Heading ([1-6])', style)
                blocks.append(('#' * int(heading.group(1)) + ' ' if heading else '') + text)
            for table in doc.tables:
                rows = [[cell.text.replace('|', '\\|').replace('\n', ' ') for cell in row.cells] for row in table.rows]
                if rows:
                    blocks.append('\n'.join(['| ' + ' | '.join(rows[0]) + ' |', '| ' + ' | '.join(['---'] * len(rows[0])) + ' |'] + ['| ' + ' | '.join(row) + ' |' for row in rows[1:]]))
            return '\n\n'.join(blocks)
        except (ValueError, KeyError, TypeError, zipfile.BadZipFile) as exc:
            raise ValidationError('DOCX 내용을 읽을 수 없습니다.') from exc
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise ValidationError('Markdown과 HTML은 UTF-8 파일이어야 합니다.') from exc
    if suffix == '.md':
        return text
    from bs4 import BeautifulSoup
    from markdownify import markdownify
    soup = BeautifulSoup(text, 'html.parser')
    for element in soup(['script', 'style', 'iframe', 'object', 'embed', 'form', 'input', 'img']):
        element.decompose()
    return markdownify(str(soup), heading_style='ATX')

@transaction.atomic
def import_documents(user, space, filename, data):
    require(user, space, 'edit')
    if len(data) > MAX_UPLOAD:
        raise ValidationError('가져오기 파일은 10MB 이하입니다.')
    documents = []
    if PurePosixPath(filename).suffix.lower() == '.zip':
        archive = validate_archive(data)
        try:
            expanded_budget = sum(item.file_size for item in archive.infolist())
            for info in archive.infolist():
                if info.is_dir():
                    continue
                if PurePosixPath(info.filename).suffix.lower() not in SUPPORTED:
                    raise ValidationError('ZIP에는 Markdown, HTML, DOCX 문서만 포함하세요.')
                payload = archive.read(info)
                if PurePosixPath(info.filename).suffix.lower() == '.docx':
                    nested = validate_archive(payload)
                    try:
                        expanded_budget += sum(item.file_size for item in nested.infolist())
                        if expanded_budget > MAX_EXPANDED:
                            raise ValidationError('DOCX를 포함한 전체 압축 해제 크기 제한을 초과했습니다.')
                    finally:
                        nested.close()
                documents.append((info.filename, decode_document(info.filename, payload)))
        except (zipfile.BadZipFile, RuntimeError) as exc:
            raise ValidationError('ZIP 파일을 읽을 수 없습니다.') from exc
        finally:
            archive.close()
    else:
        documents.append((PurePosixPath(filename).name, decode_document(filename, data)))
    if not documents:
        raise ValidationError('가져올 문서가 없습니다.')
    # Create an isolated import folder: never overwrite local administrator content or ACLs.
    root = ManualNode.objects.create(space=space, title='가져오기 · ' + PurePosixPath(filename).stem[:160], kind='folder')
    folders = {(): root}
    nodes = []
    for name, body in documents:
        path = PurePosixPath(name)
        if len(path.parts) > 20:
            raise ValidationError('가져오기 폴더는 최대 20단계입니다.')
        for depth in range(1, len(path.parts)):
            key = path.parts[:depth]
            if key not in folders:
                folders[key] = ManualNode.objects.create(space=space, parent=folders[key[:-1]], title=key[-1][:200], kind='folder')
        parent = folders[path.parts[:-1]]
        require_node(user, parent, 'create')
        node = ManualNode.objects.create(space=space, parent=parent, title=path.stem[:200] or '가져온 문서')
        save_revision(user, node.pk, node.title, body, 0, '문서 가져오기')
        node.refresh_from_db()
        nodes.append(node)
    audit(user, 'manual.import', root)
    return nodes

def published_pdf(user, node):
    require_node(user, node, 'export')
    if not node.published:
        raise ValidationError('게시된 버전만 PDF로 내보낼 수 있습니다.')
    from weasyprint import HTML
    def deny_fetch(url, *args, **kwargs):
        # HTML sanitization alone does not prevent PDF renderer SSRF or local file disclosure.
        raise ValueError('External and local resource fetching is disabled for PDF export.')
    css = '@page{size:A4;margin:20mm}body{font-family:"Noto Sans CJK KR",sans-serif;color:#19283e;font-size:11pt;line-height:1.7}h1{font-size:23pt}h2{font-size:17pt}table{border-collapse:collapse;width:100%}td,th{border:1px solid #ccd5e1;padding:6px}pre{white-space:pre-wrap;overflow-wrap:anywhere}img{display:none}'
    html = f'<!doctype html><html lang="ko"><meta charset="utf-8"><style>{css}</style><h1>{escape(node.published.title)}</h1><p>v{node.published.number}</p>{render_markdown(node.published.markdown)}</html>'
    return HTML(string=html, url_fetcher=deny_fetch).write_pdf()
